import openpyxl
import smtplib
from email.message import EmailMessage

# ==========================================
# 1. CONFIGURATION (UPDATE THESE!)
# ==========================================
EXCEL_FILE = r"C:\Users\shragaja\OneDrive - Capgemini\Documents\Vanguard_Onboarding.xlsx"
EMAIL_ADDRESS = "shraddhawarade85@gmail.com"  # Your sender email
EMAIL_PASSWORD = "" # The App Password you generated
SMTP_SERVER = "smtp.gmail.com" # Use smtp-mail.outlook.com if using Outlook
SMTP_PORT = 587

EMAIL_TEMPLATES = {
    "welcome": {
        "subject": "Welcome to Vanguard, {Name}! Action Required",
        "body": "Hi {Name},\n\nWelcome to the Vanguard project! Your account profile has been successfully created.\n\nPlease upload your BGC documents.\n\nBest,\nTeam"
    },
    "trainings": {
        "subject": "Mandatory Trainings Assigned for {Name}",
        "body": "Hi {Name},\n\nPlease complete the following 4 mandatory trainings:\n1. InfoSec\n2. Code of Conduct\n3. POSH\n4. Data Privacy\n\nBest,\nTeam"
    },
    "reminder": {
        "subject": "Action Required: Missing BGC Documents for {Name}",
        "body": "Hi {Name},\n\nYour BGC documents are still pending. Please submit them.\n\nBest,\nTeam"
    },
    "training_reminder": {
        "subject": "Reminder: Incomplete Trainings for {Name}",
        "body": "Hi {Name},\n\nYou still have these trainings pending:\n\n{PendingList}\n\nPlease complete them immediately.\n\nBest,\nTeam"
    },
    "credentials": {
        "subject": "Onboarding Complete! Your Credentials",
        "body": "Hi {Name},\n\nCongratulations, your onboarding is complete!\n\nWelcome aboard!\nTeam"
    }
}

# ==========================================
# 2. EMAIL SENDER FUNCTION
# ==========================================
def send_email(to_email, name, template_key, pending_list=None):
    template = EMAIL_TEMPLATES[template_key]
    subject = template["subject"].replace("{Name}", name)
    body = template["body"].replace("{Name}", name)
    
    if pending_list:
        body = body.replace("{PendingList}", "\n".join(pending_list))
        
    msg = EmailMessage()
    msg.set_content(body)
    msg['Subject'] = subject
    msg['From'] = EMAIL_ADDRESS
    msg['To'] = to_email
    
    try:
        server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
        server.starttls()
        server.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
        server.send_message(msg)
        server.quit()
        print(f"✅ Sent '{template_key}' email to {name} ({to_email})")
        return True
    except Exception as e:
        print(f"❌ Failed to send email to {to_email}. Error: {e}")
        return False

# ==========================================
# 3. AUTOMATION ENGINE
# ==========================================
def run_automations():
    print("Starting Vanguard Onboarding Automation...")
    wb = openpyxl.load_workbook(EXCEL_FILE)
    sheet = wb.active
    
    # openpyxl starts at row 1. Data starts at row 2.
    for i in range(2, sheet.max_row + 1):
        name = sheet.cell(row=i, column=1).value
        email = sheet.cell(row=i, column=2).value
        
        if not name or not email:
            continue
            
        profile_created = sheet.cell(row=i, column=3).value
        bgc_status = sheet.cell(row=i, column=4).value
        onboarding_status = sheet.cell(row=i, column=5).value
        
        welcome_sent = sheet.cell(row=i, column=6).value
        reminder_sent = sheet.cell(row=i, column=7).value
        trainings_sent = sheet.cell(row=i, column=8).value
        credentials_sent = sheet.cell(row=i, column=9).value
        
        info_sec = sheet.cell(row=i, column=10).value
        code_of_conduct = sheet.cell(row=i, column=11).value
        posh = sheet.cell(row=i, column=12).value
        data_privacy = sheet.cell(row=i, column=13).value
        training_reminder_sent = sheet.cell(row=i, column=14).value
        
        # TRIGGER 1: Welcome Email
        if profile_created == "Yes" and welcome_sent != "Yes":
            if send_email(email, name, "welcome"):
                sheet.cell(row=i, column=6).value = "Yes"
            continue
            
        # TRIGGER 2: Trainings Email
        if welcome_sent == "Yes" and trainings_sent != "Yes":
            if send_email(email, name, "trainings"):
                sheet.cell(row=i, column=8).value = "Yes"
            continue
            
        # TRIGGER 3: BGC Reminder
        if bgc_status == "Pending" and welcome_sent == "Yes" and reminder_sent != "Yes":
            if send_email(email, name, "reminder"):
                sheet.cell(row=i, column=7).value = "Yes"
            continue
            
        # TRIGGER 4: Dynamic Training Reminder
        if trainings_sent == "Yes" and training_reminder_sent != "Yes":
            pending_list = []
            if info_sec == "Pending" or not info_sec: pending_list.append("- Information Security")
            if code_of_conduct == "Pending" or not code_of_conduct: pending_list.append("- Code of Conduct")
            if posh == "Pending" or not posh: pending_list.append("- POSH")
            if data_privacy == "Pending" or not data_privacy: pending_list.append("- Data Privacy")
            
            if pending_list:
                if send_email(email, name, "training_reminder", pending_list):
                    sheet.cell(row=i, column=14).value = "Yes"
                continue
                
        # TRIGGER 5: Credentials
        if onboarding_status == "Completed" and credentials_sent != "Yes":
            if send_email(email, name, "credentials"):
                sheet.cell(row=i, column=9).value = "Yes"
            continue

    # Save the Excel file to lock in the "Yes" updates
    wb.save(EXCEL_FILE)
    print("Automation complete! Excel file updated.")

if __name__ == "__main__":
    run_automations()
